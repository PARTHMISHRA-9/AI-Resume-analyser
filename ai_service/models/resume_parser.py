import PyPDF2
import docx
import re
from typing import Dict, List

class ResumeParser:
    """Parse resume from PDF or Word documents"""

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract text from PDF file"""
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text()
        except Exception as e:
            raise Exception(f"Error parsing PDF: {str(e)}")
        return text

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from Word document"""
        text = ""
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
        except Exception as e:
            raise Exception(f"Error parsing DOCX: {str(e)}")
        return text

    @staticmethod
    def extract_text(file_path: str, file_type: str) -> str:
        """Extract text based on file type"""
        if file_type == 'pdf':
            return ResumeParser.extract_text_from_pdf(file_path)
        elif file_type in ['doc', 'docx']:
            return ResumeParser.extract_text_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    @staticmethod
    def extract_emails(text: str) -> List[str]:
        """Extract email addresses from text"""
        pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
        return re.findall(pattern, text)

    @staticmethod
    def extract_phone_numbers(text: str) -> List[str]:
        """Extract phone numbers from text"""
        pattern = r'\b(?:\d{3}[-.]?\d{3}[-.]?\d{4}|\d{10})\b'
        return re.findall(pattern, text)

    @staticmethod
    def extract_sections(text: str) -> Dict[str, str]:
        """Extract common resume sections"""
        sections = {}
        section_keywords = {
            'summary': ['summary', 'objective', 'professional summary'],
            'experience': ['experience', 'work history', 'employment'],
            'education': ['education', 'academic', 'qualification'],
            'skills': ['skills', 'technical skills', 'competencies'],
            'projects': ['projects', 'portfolio', 'achievements'],
            'certifications': ['certifications', 'licenses', 'awards'],
        }

        lines = text.split('\n')
        current_section = None
        section_content = []

        for line in lines:
            line_lower = line.lower().strip()
            matched = False

            for section, keywords in section_keywords.items():
                if any(keyword in line_lower for keyword in keywords):
                    if current_section:
                        sections[current_section] = '\n'.join(section_content)
                    current_section = section
                    section_content = []
                    matched = True
                    break

            if not matched and current_section:
                section_content.append(line)

        if current_section:
            sections[current_section] = '\n'.join(section_content)

        return sections
